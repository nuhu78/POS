from django.test import TestCase
from rest_framework.test import APIClient
import io
from openpyxl import load_workbook
from apps.accounts.models import User
from apps.categories.models import Category
from .models import Product


class ProductPermissionTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_user(
            email="admin@test.com", name="Admin", password="pass123", role="admin"
        )
        self.cashier = User.objects.create_user(
            email="cashier@test.com", name="Cashier", password="pass123", role="cashier"
        )
        self.cat = Category.objects.create(name="Furniture")

    def test_admin_can_create_product(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.post("/api/v1/products/", {
            "name": "Chair", "sku": "CHR-001", "category": self.cat.id,
            "purchase_price": "50.00", "selling_price": "120.00",
            "stock": 10, "status": "active",
        })
        self.assertEqual(resp.status_code, 201)

    def test_cashier_cannot_create_product(self):
        self.client.force_authenticate(user=self.cashier)
        resp = self.client.post("/api/v1/products/", {
            "name": "Chair", "sku": "CHR-001", "category": self.cat.id,
            "purchase_price": "50.00", "selling_price": "120.00",
            "stock": 10, "status": "active",
        })
        self.assertEqual(resp.status_code, 403)

    def test_cashier_can_list_products(self):
        self.client.force_authenticate(user=self.cashier)
        resp = self.client.get("/api/v1/products/")
        self.assertEqual(resp.status_code, 200)

    def test_admin_can_delete_product(self):
        self.client.force_authenticate(user=self.admin)
        p = Product.objects.create(
            name="Chair", sku="CHR-001", category=self.cat,
            purchase_price=50, selling_price=120, stock=10,
        )
        resp = self.client.delete(f"/api/v1/products/{p.id}/")
        self.assertEqual(resp.status_code, 204)

    def test_cashier_cannot_delete_product(self):
        self.client.force_authenticate(user=self.admin)
        p = Product.objects.create(
            name="Chair", sku="CHR-001", category=self.cat,
            purchase_price=50, selling_price=120, stock=10,
        )
        self.client.force_authenticate(user=self.cashier)
        resp = self.client.delete(f"/api/v1/products/{p.id}/")
        self.assertEqual(resp.status_code, 403)


class ProductFilterTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_user(
            email="admin@test.com", name="Admin", password="pass123", role="admin"
        )
        self.cat_a = Category.objects.create(name="Furniture")
        self.cat_b = Category.objects.create(name="Gifts")
        Product.objects.create(
            name="Chair", sku="CHR-001", category=self.cat_a,
            purchase_price=50, selling_price=120, stock=10, low_stock_threshold=5,
        )
        Product.objects.create(
            name="Low Chair", sku="CHR-002", category=self.cat_a,
            purchase_price=40, selling_price=90, stock=2, low_stock_threshold=5,
        )
        Product.objects.create(
            name="Vase", sku="VAS-001", category=self.cat_b,
            purchase_price=10, selling_price=30, stock=0, low_stock_threshold=5,
        )
        Product.objects.create(
            name="Old Vase", sku="VAS-002", category=self.cat_b,
            purchase_price=10, selling_price=25, stock=0, low_stock_threshold=5,
            status="inactive",
        )
        self.client.force_authenticate(user=self.admin)

    def skus(self, resp):
        return [r["sku"] for r in resp.data["results"]]

    def test_filter_by_price_range(self):
        resp = self.client.get("/api/v1/products/", {"price_min": 100, "price_max": 130})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(self.skus(resp), ["CHR-001"])

    def test_filter_by_price_min_only(self):
        resp = self.client.get("/api/v1/products/", {"price_min": 100})
        self.assertEqual(self.skus(resp), ["CHR-001"])

    def test_filter_low_stock(self):
        resp = self.client.get("/api/v1/products/", {"stock_status": "low"})
        self.assertEqual(sorted(self.skus(resp)), ["CHR-002", "VAS-001"])

    def test_filter_out_of_stock(self):
        resp = self.client.get("/api/v1/products/", {"stock_status": "out"})
        self.assertEqual(sorted(self.skus(resp)), ["VAS-001", "VAS-002"])

    def test_filter_by_category(self):
        resp = self.client.get("/api/v1/products/", {"category": self.cat_a.id})
        self.assertEqual(sorted(self.skus(resp)), ["CHR-001", "CHR-002"])

    def test_combined_category_and_price(self):
        resp = self.client.get("/api/v1/products/", {
            "category": self.cat_a.id, "price_min": 100, "price_max": 130,
        })
        self.assertEqual(self.skus(resp), ["CHR-001"])

    def exported_skus(self, params):
        resp = self.client.get("/api/v1/products/export/", params)
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content), read_only=True)
        ws = wb.active
        return [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]

    def test_export_respects_stock_filter(self):
        self.assertEqual(
            sorted(self.exported_skus({"stock_status": "out"})),
            ["VAS-001", "VAS-002"],
        )

    def test_export_respects_price_filter(self):
        self.assertEqual(
            self.exported_skus({"price_min": 100, "price_max": 130}),
            ["CHR-001"],
        )

    def test_export_respects_search(self):
        self.assertEqual(
            sorted(self.exported_skus({"search": "chair"})),
            ["CHR-001", "CHR-002"],
        )

    def test_export_respects_category_filter(self):
        self.assertEqual(
            sorted(self.exported_skus({"category": self.cat_a.id})),
            ["CHR-001", "CHR-002"],
        )
